import io
import pytest
from PIL import Image, ImageDraw
from fastapi.testclient import TestClient
from app.main import app
from app.services.report_service import get_report_service

client = TestClient(app)

def create_dummy_png_bytes() -> bytes:
    img = Image.new("RGB", (300, 200), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((20, 20), "PackTrue Sample Label", fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()

@pytest.fixture
def sample_scan_record():
    img_bytes = create_dummy_png_bytes()
    scan_id = "testscan123"
    scan_data = {
        "scan_id": scan_id,
        "filename": "sample_commodity.png",
        "image_bytes": img_bytes,
        "timestamp": "2026-09-09 12:00:00 UTC",
        "scan_result": {
            "success": True,
            "filename": "sample_commodity.png",
            "image_width": 300,
            "image_height": 200,
            "total_lines": 5,
            "average_confidence": 0.95,
            "lines": [],
            "raw_text": "Sample Commodity Text",
            "processing_time_ms": 120.5,
            "ocr_engine": "PaddleOCR",
            "image_quality": {
                "blur_score": 250.0,
                "quality_threshold": 100.0,
                "is_blurry": False,
                "enhancement_applied": False,
                "quality_status": "Good — no enhancement needed",
                "enhancement_method": "None"
            },
            "compliance": {
                "overall_verdict": "Compliant",
                "summary": "Label conforms to Legal Metrology Rule 6 requirements.",
                "total_checks": 6,
                "passed_checks": 6,
                "advisory_checks": 0,
                "failed_checks": 0,
                "fields": [
                    {
                        "field_id": "manufacturer",
                        "field_name": "Manufacturer / Packer Details",
                        "extracted_value": "Pure Dairy Ltd, Anand 388001",
                        "verdict": "Compliant",
                        "reason": "Name and complete postal address with PIN code present.",
                        "rule_reference": "Rule 6(1)(a)"
                    },
                    {
                        "field_id": "net_quantity",
                        "field_name": "Net Quantity",
                        "extracted_value": "1.0 L (905 g)",
                        "verdict": "Compliant",
                        "reason": "Standard SI metric units specified.",
                        "rule_reference": "Rule 6(1)(b) & Rule 12"
                    },
                    {
                        "field_id": "mrp",
                        "field_name": "Maximum Retail Price (MRP)",
                        "extracted_value": "Rs. 590.00 (INCL. OF ALL TAXES)",
                        "verdict": "Compliant",
                        "reason": "MRP with mandatory tax-inclusive declaration present.",
                        "rule_reference": "Rule 6(1)(e)"
                    },
                    {
                        "field_id": "mfg_date",
                        "field_name": "Month & Year of Manufacture",
                        "extracted_value": "12/04/2026",
                        "verdict": "Compliant",
                        "reason": "Unambiguous manufacturing date present.",
                        "rule_reference": "Rule 6(1)(d)"
                    },
                    {
                        "field_id": "consumer_care",
                        "field_name": "Consumer Care Details",
                        "extracted_value": "1800-258-3333 / care@puredairy.in",
                        "verdict": "Compliant",
                        "reason": "Helpline and email address verified.",
                        "rule_reference": "Rule 6(1)(f)"
                    },
                    {
                        "field_id": "fssai_license",
                        "field_name": "FSSAI License Declaration",
                        "extracted_value": "10014021000001",
                        "verdict": "Compliant",
                        "reason": "14-digit statutory license present.",
                        "rule_reference": "FSS Regulations"
                    }
                ]
            }
        }
    }
    svc = get_report_service()
    svc.save_scan(
        scan_id=scan_id,
        filename=scan_data["filename"],
        image_bytes=scan_data["image_bytes"],
        scan_result=scan_data["scan_result"],
        timestamp=scan_data["timestamp"]
    )
    return scan_data


def test_generate_pdf_report(sample_scan_record):
    svc = get_report_service()
    pdf_bytes = svc.generate_pdf_report(sample_scan_record)
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 500
    # PDF magic bytes
    assert pdf_bytes.startswith(b"%PDF-")


def test_generate_docx_report(sample_scan_record):
    svc = get_report_service()
    docx_bytes = svc.generate_docx_report(sample_scan_record)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 500
    # DOCX is a zip file, magic bytes PK\x03\x04
    assert docx_bytes.startswith(b"PK\x03\x04")


def test_get_report_pdf_endpoint(sample_scan_record):
    scan_id = sample_scan_record["scan_id"]
    response = client.get(f"/api/v1/report/{scan_id}?format=pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith(f"packtrue_compliance_report_{scan_id}.pdf\"")
    assert response.content.startswith(b"%PDF-")


def test_get_report_docx_endpoint(sample_scan_record):
    scan_id = sample_scan_record["scan_id"]
    response = client.get(f"/api/v1/report/{scan_id}?format=docx")
    assert response.status_code == 200
    assert "officedocument.wordprocessingml.document" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith(f"packtrue_compliance_report_{scan_id}.docx\"")
    assert response.content.startswith(b"PK\x03\x04")


def test_get_report_not_found():
    response = client.get("/api/v1/report/nonexistent999?format=pdf")
    assert response.status_code == 404
    assert "not found" in response.json()["detail"].lower()


def test_get_report_invalid_format(sample_scan_record):
    scan_id = sample_scan_record["scan_id"]
    response = client.get(f"/api/v1/report/{scan_id}?format=exe")
    assert response.status_code == 400
    assert "unsupported format" in response.json()["detail"].lower()


def test_report_preliminary_assessment_wording(sample_scan_record):
    """Verify that reports feature the required AI-assisted pre-screening framing and disclaimer."""
    import docx
    svc = get_report_service()
    docx_bytes = svc.generate_docx_report(sample_scan_record)

    doc = docx.Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += " " + cell.text

    # 1. Status reworded
    assert "AI-Assisted Pre-Screening — Requires Officer Confirmation" in all_text
    assert "Official Statutory Audit" not in all_text

    # 2. Automated pre-screening disclaimer
    assert "This is an automated pre-screening tool" in all_text
    assert "Findings are not a legal determination and require confirmation" in all_text

    # 3. Assessment title
    assert "PRELIMINARY ASSESSMENT" in all_text


def test_report_non_compliant_potential_violations_wording(sample_scan_record):
    """Verify that non-compliant reports use 'PRELIMINARY ASSESSMENT: Potential Violations Flagged'."""
    import docx
    import copy
    rec = copy.deepcopy(sample_scan_record)
    rec["scan_result"]["compliance"]["overall_verdict"] = "Non-Compliant"

    svc = get_report_service()
    docx_bytes = svc.generate_docx_report(rec)

    doc = docx.Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += " " + cell.text

    assert "PRELIMINARY ASSESSMENT: Potential Violations Flagged" in all_text
    assert "AUDIT VERDICT: NON-COMPLIANT" not in all_text


def test_report_low_confidence_banner(sample_scan_record):
    """Verify that if image quality was low (enhancement applied), the low confidence review banner is shown."""
    import docx
    import copy
    rec = copy.deepcopy(sample_scan_record)
    rec["scan_result"]["image_quality"]["enhancement_applied"] = True
    rec["scan_result"]["image_quality"]["is_blurry"] = True

    svc = get_report_service()
    docx_bytes = svc.generate_docx_report(rec)

    doc = docx.Document(io.BytesIO(docx_bytes))
    all_text = " ".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_text += " " + cell.text

    assert "Low Confidence — Human Review Recommended" in all_text

